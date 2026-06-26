import logging
import re
from typing import Any, Dict, List, Optional

from ..core.definitive_style_markers import (
    FW_BLOCK_END,
    FW_EPISTEMIC_BEGIN,
    FW_EPISTEMIC_END,
    RE_FW_BLOCK_OPEN,
    RE_FW_HEADING,
)
from ..core.mechanical_audit import definitive_document_for_export

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches
except Exception:  # pragma: no cover
    Document = None
    WD_ALIGN_PARAGRAPH = None
    Pt = None
    RGBColor = None
    Inches = None
    OxmlElement = None
    qn = None

log = logging.getLogger(__name__)

_DOCX_FONT = "Calibri"
_DOCX_EPISTEMIC_INDENT_PT = 18
# List: same left gap as epistemic; hang = bullet column so wrapped lines align with text.
_DOCX_LIST_INDENT_PT = 18
_DOCX_LIST_HANG_PT = 18
_TABLE_GAP_AFTER_PT = 6
_CELL_SPACE_AFTER_PT = 4

_DOCX_SUBSECTION_RGB = {
    "definition": (13, 71, 161),
    "principle": (13, 71, 161),
    "step": (17, 17, 17),
    "warning": (230, 81, 0),
    "constraint": (27, 94, 32),
}


def _run_font(run, size_pt: Optional[Any] = None) -> None:
    run.font.name = _DOCX_FONT
    if Pt is not None and size_pt is not None:
        run.font.size = size_pt


def _apply_table_gap_before(ctx: Dict[str, Any], p) -> None:
    gap = ctx.get("pad_before_next_block")
    if gap is not None and Pt is not None:
        p.paragraph_format.space_before = gap
    ctx["pad_before_next_block"] = None


def _mark_table_ended(ctx: Dict[str, Any]) -> None:
    if Pt is not None:
        ctx["pad_before_next_block"] = Pt(_TABLE_GAP_AFTER_PT)


def _cell_para_spacing(p) -> None:
    if Pt is None:
        return
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(_CELL_SPACE_AFTER_PT)


def _docx_append_bold_segments(
    p,
    text: str,
    font_size: Optional[Any] = None,
    color_rgb: Optional[Any] = None,
) -> None:
    segments = re.split(r"(\*\*.+?\*\*)", text)
    for seg in segments:
        if not seg:
            continue
        seg_stripped = seg.strip()
        bold_m = re.match(r"^\*\*(.+?)\*\*$", seg_stripped)
        run = p.add_run(bold_m.group(1) if bold_m else seg)
        run.bold = bool(bold_m)
        _run_font(run, font_size)
        if RGBColor is not None and color_rgb is not None:
            run.font.color.rgb = color_rgb


def _docx_append_markdown_inline(
    p,
    text: str,
    font_size: Optional[Any] = None,
    color_rgb: Optional[Any] = None,
) -> None:
    """**bold** and _italic_ (same idea as UI); plain body lines and bullets."""
    if not text:
        return
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        seg_stripped = seg.strip()
        bold_m = re.match(r"^\*\*(.+?)\*\*$", seg_stripped)
        if bold_m:
            run = p.add_run(bold_m.group(1))
            run.bold = True
            _run_font(run, font_size)
            if RGBColor is not None and color_rgb is not None:
                run.font.color.rgb = color_rgb
            continue
        for piece in re.split(r"(_[^_\n]+?_)", seg):
            if not piece:
                continue
            im = re.match(r"^_(.+?)_$", piece)
            if im:
                run = p.add_run(im.group(1))
                run.italic = True
                _run_font(run, font_size)
                if RGBColor is not None and color_rgb is not None:
                    run.font.color.rgb = color_rgb
            else:
                run = p.add_run(piece)
                _run_font(run, font_size)
                if RGBColor is not None and color_rgb is not None:
                    run.font.color.rgb = color_rgb


def _split_source_epistemic_label_body(s: str) -> tuple[str, str]:
    """Strip markdown wrappers; return (label through ':', quote body after ':')."""
    raw = (s or "").strip()
    m = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", raw, re.DOTALL)
    if m:
        t, tail = m.group(1).strip(), m.group(2).strip()
    else:
        t = raw
        if t.startswith("**"):
            t = t[2:].strip()
            if t.endswith("**"):
                t = t[:-2].strip()
        tail = ""
    t = re.sub(r"^_+", "", t)
    ix = t.find(":")
    if ix < 0:
        lab = t.strip("_ \t")
        bod = tail
        return lab, bod.strip()
    lab = t[: ix + 1].strip().strip("_ \t")
    bod = t[ix + 1 :].lstrip()
    bod = re.sub(r"^_+\s*", "", bod)
    bod = bod.strip()
    if tail:
        bod = f"{bod} {tail}".strip() if bod else tail
    return lab, bod


def _append_fw_epistemic_docx(doc: "Document", ep_lines: List[str], ctx: Optional[Dict[str, Any]] = None) -> None:
    if Document is None:
        return
    small = Pt(11) if Pt is not None else None
    ind = Pt(_DOCX_EPISTEMIC_INDENT_PT) if Pt is not None else None
    purple = RGBColor(106, 27, 154) if RGBColor is not None else None
    purple_body = RGBColor(74, 20, 140) if RGBColor is not None else None
    blue_label = RGBColor(25, 118, 210) if RGBColor is not None else None
    dark_body = RGBColor(26, 26, 26) if RGBColor is not None else None
    first_epi = True

    def _epi_para() -> Any:
        nonlocal first_epi
        p = doc.add_paragraph()
        if first_epi:
            if ctx is not None:
                gap = ctx.get("pad_before_next_block")
                _apply_table_gap_before(ctx, p)
                if Pt is not None and gap is None:
                    p.paragraph_format.space_before = Pt(0)
            elif Pt is not None:
                p.paragraph_format.space_before = Pt(0)
            first_epi = False
        elif Pt is not None:
            p.paragraph_format.space_before = Pt(0)
        if Pt is not None:
            p.paragraph_format.left_indent = ind
            p.paragraph_format.space_after = Pt(2)
        return p

    for raw in ep_lines:
        s = (raw or "").strip()
        if not s:
            continue

        if "**Tag justifications:**" in s:
            ix = s.find("**Tag justifications:**")
            prefix = s[:ix].strip()
            rest = s[ix + len("**Tag justifications:**") :].lstrip()
            p = _epi_para()
            if prefix:
                _docx_append_bold_segments(p, prefix, font_size=small, color_rgb=dark_body)
                p.add_run(" ")
            r_lab = p.add_run("Tag justifications:")
            r_lab.bold = True
            _run_font(r_lab, small)
            if purple is not None:
                r_lab.font.color.rgb = purple
            if rest:
                p.add_run(" ")
                _docx_append_bold_segments(p, rest, font_size=small, color_rgb=dark_body)
            continue

        if s.startswith("**ES:**"):
            p = _epi_para()
            _docx_set_paragraph_shading(p, "EDE7F6")
            _docx_append_bold_segments(p, s, font_size=small, color_rgb=purple_body)
            continue

        if s.startswith("**_Source") or s.startswith("**Source"):
            p = _epi_para()
            _docx_set_paragraph_shading(p, "F5F5F5")
            label, body = _split_source_epistemic_label_body(s)
            r_lab = p.add_run(label)
            r_lab.bold = True
            _run_font(r_lab, small)
            if blue_label is not None:
                r_lab.font.color.rgb = blue_label
            if body:
                p.add_run(" ")
                _docx_append_bold_segments(p, body, font_size=small, color_rgb=dark_body)
            continue

        p = _epi_para()
        _docx_append_bold_segments(p, s, font_size=small, color_rgb=dark_body)


def _docx_set_cell_shading(cell, fill_hex: str) -> None:
    if OxmlElement is None or qn is None:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _fill_docx_table_cell(
    cell,
    clean_text: str,
    *,
    bold: bool,
    shade_hex: Optional[str] = None,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    _cell_para_spacing(p)
    run = p.add_run(clean_text)
    run.bold = bold
    if Pt is not None:
        _run_font(run, Pt(12))
    if shade_hex:
        _docx_set_cell_shading(cell, shade_hex)


def _docx_add_markdown_bullet(doc: "Document", ctx: Dict[str, Any], bullet_md: str) -> None:
    """Manual bullet + hanging indent (List Bullet style fights indents; matches epistemic left gap)."""
    if Document is None or Pt is None:
        return
    g = Pt(_DOCX_LIST_INDENT_PT)
    h = Pt(_DOCX_LIST_HANG_PT)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = g + h
    p.paragraph_format.first_line_indent = -h
    _run_font(p.add_run("• "), Pt(12))
    _apply_table_gap_before(ctx, p)
    p.paragraph_format.space_after = Pt(2)
    _docx_append_markdown_inline(p, bullet_md, font_size=Pt(12), color_rgb=None)


def _docx_set_paragraph_shading(paragraph, fill_hex: str) -> None:
    if OxmlElement is None or qn is None:
        return
    p_el = getattr(paragraph, "_p", None) or getattr(paragraph, "_element", None)
    if p_el is None:
        return
    p_pr = p_el.get_or_add_pPr()
    for el in list(p_pr):
        if el.tag == qn("w:shd"):
            p_pr.remove(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill_hex)
    p_pr.append(shd)


def _docx_set_table_full_page_width(table) -> None:
    if OxmlElement is None or qn is None:
        return
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    for el in list(tbl_pr):
        if el.tag == qn("w:tblW"):
            tbl_pr.remove(el)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")
    tbl_pr.append(tbl_w)


def _docx_apply_normal_page_margins(doc: "Document") -> None:
    if Inches is None:
        return
    inch1 = Inches(1)
    inch_bottom = Inches(1.15)
    for sec in doc.sections:
        sec.left_margin = inch1
        sec.right_margin = inch1
        sec.top_margin = inch1
        sec.bottom_margin = inch_bottom


def _docx_apply_document_default_font(doc: "Document") -> None:
    if Pt is None:
        return
    try:
        n = doc.styles["Normal"].font
        n.name = _DOCX_FONT
        n.size = Pt(12)
    except (KeyError, AttributeError):
        pass


def _docx_append_complex_field(paragraph, instruction: str, placeholder: str = "0") -> None:
    if OxmlElement is None or qn is None:
        return
    run = paragraph.add_run()
    _run_font(run, Pt(9) if Pt is not None else None)
    r = run._r
    fc_b = OxmlElement("w:fldChar")
    fc_b.set(qn("w:fldCharType"), "begin")
    r.append(fc_b)
    ins = OxmlElement("w:instrText")
    ins.set(qn("xml:space"), "preserve")
    ins.text = f" {instruction} "
    r.append(ins)
    fc_s = OxmlElement("w:fldChar")
    fc_s.set(qn("w:fldCharType"), "separate")
    r.append(fc_s)
    t = OxmlElement("w:t")
    t.text = placeholder
    r.append(t)
    fc_e = OxmlElement("w:fldChar")
    fc_e.set(qn("w:fldCharType"), "end")
    r.append(fc_e)


def _docx_add_footer_page_numbers(doc: "Document") -> None:
    if Document is None or WD_ALIGN_PARAGRAPH is None or OxmlElement is None:
        return
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p.add_run("Page ")
        _run_font(r0, Pt(9) if Pt is not None else None)
        _docx_append_complex_field(p, "PAGE")
        r1 = p.add_run(" of ")
        _run_font(r1, Pt(9) if Pt is not None else None)
        _docx_append_complex_field(p, "NUMPAGES")


def _add_docx_heading(doc: "Document", text: str, ctx: Dict[str, Any]) -> None:
    p = doc.add_paragraph()
    _apply_table_gap_before(ctx, p)
    run = p.add_run(text)
    run.bold = True
    if Pt is not None:
        _run_font(run, Pt(16))
        p.paragraph_format.space_after = Pt(3)


def _append_structure_to_docx(doc: "Document", text: str) -> None:
    lines = text.splitlines()
    i = 0
    ctx: Dict[str, Any] = {
        "pending_heading": None,
        "fw_block_depth": 0,
        "pad_before_next_block": None,
    }
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip("\n")
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        m_block = RE_FW_BLOCK_OPEN.match(stripped)
        if m_block:
            ctx["fw_block_depth"] = int(ctx.get("fw_block_depth") or 0) + 1
            i += 1
            continue
        if stripped == FW_BLOCK_END:
            d = int(ctx.get("fw_block_depth") or 0)
            if d > 0:
                ctx["fw_block_depth"] = d - 1
            i += 1
            continue
        mh = RE_FW_HEADING.match(stripped)
        if mh:
            ctx["pending_heading"] = mh.group(1).lower()
            i += 1
            continue
        if stripped == FW_EPISTEMIC_BEGIN:
            j = i + 1
            acc: List[str] = []
            while j < len(lines):
                sl = lines[j].strip()
                if sl == FW_EPISTEMIC_END:
                    j += 1
                    break
                acc.append(lines[j])
                j += 1
            _append_fw_epistemic_docx(doc, acc, ctx)
            i = j
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            _apply_table_gap_before(ctx, p)
            if Pt is not None:
                p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if stripped.startswith(">"):
            quote_lines: list[str] = []
            j = i
            while j < len(lines):
                s = lines[j].strip()
                if s.startswith(">"):
                    quote_lines.append(s[1:].lstrip())
                    j += 1
                elif not s:
                    j += 1
                    break
                else:
                    break
            first_q = True
            for ql in quote_lines:
                if not ql:
                    continue
                p = doc.add_paragraph()
                if first_q:
                    _apply_table_gap_before(ctx, p)
                    first_q = False
                if Pt is not None:
                    p.paragraph_format.left_indent = Pt(18)
                    p.paragraph_format.space_after = Pt(2)
                if "**" in ql:
                    segments = re.split(r"(\*\*.+?\*\*)", ql)
                    for seg in segments:
                        if not seg:
                            continue
                        m = re.match(r"^\*\*(.+?)\*\*$", seg.strip())
                        if m:
                            r = p.add_run(m.group(1))
                            r.bold = True
                            _run_font(r, Pt(12))
                        else:
                            r = p.add_run(seg)
                            _run_font(r, Pt(12))
                else:
                    r = p.add_run(ql)
                    _run_font(r, Pt(12))
            i = j
            continue

        if stripped.startswith("|") and (i + 1) < len(lines):
            nxt = lines[i + 1].strip()
            if nxt.startswith("|") and "---" in nxt:
                table_lines: list[str] = []
                j = i
                while j < len(lines) and lines[j].strip().startswith("|"):
                    table_lines.append(lines[j].strip())
                    j += 1

                def split_row(r: str) -> list[str]:
                    r = r.strip()
                    if r.startswith("|"):
                        r = r[1:]
                    if r.endswith("|"):
                        r = r[:-1]
                    return [c.strip() for c in r.split("|")]

                header_cells = split_row(table_lines[0]) if table_lines else []
                data_rows = table_lines[2:] if len(table_lines) >= 3 else []
                cols = len(header_cells) if header_cells else 1

                def _strip_markdown_bold(cell_text: str) -> tuple[str, bool]:
                    s = (cell_text or "").strip()
                    m = re.match(r"^\*\*(.+?)\*\*$", s)
                    if m:
                        return m.group(1), True
                    return s, False

                parsed_rows = [header_cells] + [split_row(dr) for dr in data_rows if dr.strip()]
                two_col_ok = cols == 2 and all(len(r) == 2 for r in parsed_rows)

                if two_col_ok:
                    table = doc.add_table(rows=len(parsed_rows), cols=2)
                    table.style = "Table Grid"
                    _docx_set_table_full_page_width(table)
                    for ridx, row_cells in enumerate(parsed_rows):
                        for c in range(2):
                            raw_cell = row_cells[c] if c < len(row_cells) else ""
                            clean, _ = _strip_markdown_bold(raw_cell)
                            cell = table.rows[ridx].cells[c]
                            _fill_docx_table_cell(
                                cell,
                                clean,
                                bold=(c == 0),
                                shade_hex="E8E8E8" if c == 0 else None,
                            )
                else:
                    table = doc.add_table(rows=len(data_rows) + 1, cols=cols)
                    table.style = "Table Grid"
                    _docx_set_table_full_page_width(table)
                    for c in range(cols):
                        cell_text = header_cells[c] if c < len(header_cells) else ""
                        clean, _ = _strip_markdown_bold(cell_text)
                        cell = table.rows[0].cells[c]
                        _fill_docx_table_cell(
                            cell, clean, bold=True, shade_hex="E8E8E8"
                        )

                    for ridx, dr in enumerate(data_rows, start=1):
                        cells = split_row(dr)
                        for c in range(cols):
                            raw_cell = cells[c] if c < len(cells) else ""
                            clean, is_bold_mark = _strip_markdown_bold(raw_cell)
                            cell = table.rows[ridx].cells[c]
                            _fill_docx_table_cell(
                                cell,
                                clean,
                                bold=is_bold_mark or (c == 0),
                                shade_hex=None,
                            )

                _mark_table_ended(ctx)
                i = j
                continue

        if stripped.startswith("## "):
            heading_text = stripped[3:].strip()
            _add_docx_heading(doc, heading_text, ctx)
            i += 1
            continue

        if stripped.startswith("### "):
            heading_text = stripped[4:].strip()
            p = doc.add_paragraph()
            _apply_table_gap_before(ctx, p)
            r = p.add_run(heading_text)
            r.bold = True
            if Pt is not None:
                _run_font(r, Pt(13))
                p.paragraph_format.space_after = Pt(2)
            if RGBColor is not None:
                r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if (
            stripped.startswith("- ")
            or stripped.startswith("* ")
            or stripped.startswith("+ ")
        ):
            _docx_add_markdown_bullet(doc, ctx, stripped[2:].strip())
            i += 1
            continue

        m = re.match(r"\*\*(.+?)\*\*(.*)", stripped)
        if m:
            head = m.group(1).strip()
            tail = m.group(2)
            is_context = head == "Context:"
            pending_k = ctx.get("pending_heading")
            sub_rgb = None
            if pending_k and not is_context and RGBColor is not None:
                sub_rgb = _DOCX_SUBSECTION_RGB.get(str(pending_k).lower(), (17, 17, 17))
                ctx["pending_heading"] = None

            p = doc.add_paragraph()
            _apply_table_gap_before(ctx, p)
            r1 = p.add_run(head)
            r1.bold = True
            if Pt is not None:
                sz = Pt(12) if is_context else (Pt(14) if sub_rgb is not None else Pt(12))
                _run_font(r1, sz)
                p.paragraph_format.space_after = Pt(2)
            if is_context and RGBColor is not None:
                r1.font.color.rgb = RGBColor(0, 0, 0)
            elif sub_rgb is not None and RGBColor is not None:
                r1.font.color.rgb = RGBColor(*sub_rgb)

            if tail:
                sz = Pt(12) if is_context else (Pt(14) if sub_rgb is not None else Pt(12))
                tail_rgb = None
                if is_context and RGBColor is not None:
                    tail_rgb = RGBColor(0, 0, 0)
                elif sub_rgb is not None and RGBColor is not None:
                    tail_rgb = RGBColor(*sub_rgb)
                _docx_append_markdown_inline(p, tail, sz, tail_rgb)
            i += 1
            continue

        if "**" in stripped:
            p = doc.add_paragraph()
            _apply_table_gap_before(ctx, p)
            _docx_append_markdown_inline(p, stripped, Pt(12), None)
        else:
            p = doc.add_paragraph()
            _apply_table_gap_before(ctx, p)
            _docx_append_markdown_inline(p, stripped, Pt(12), None)
        if Pt is not None:
            p.paragraph_format.space_after = Pt(2)
        i += 1

    log.debug("DOCX export: processed %s lines", len(lines))


def export_framework_to_docx_file(framework: Dict[str, Any], path: str) -> None:
    if Document is None or WD_ALIGN_PARAGRAPH is None:
        raise RuntimeError("python-docx is not installed. Install python-docx for DOCX export.")

    name = (framework.get("name") or "Framework").strip() or "Framework"
    definitive = definitive_document_for_export(framework).strip()
    if not definitive:
        raise RuntimeError(
            "This framework has no definitive document. Generate or regenerate before exporting to DOCX."
        )

    log.info("DOCX export: name=%s path=%s len=%s", name, path, len(definitive))

    doc = Document()
    _docx_apply_document_default_font(doc)
    _docx_apply_normal_page_margins(doc)

    title_p = doc.add_paragraph()
    run = title_p.add_run(name)
    run.bold = True
    if Pt is not None:
        _run_font(run, Pt(16))
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _append_structure_to_docx(doc, definitive)
    _docx_add_footer_page_numbers(doc)

    doc.save(path)
    log.info("DOCX export complete: %s", path)


export_framework_to_docx = export_framework_to_docx_file
